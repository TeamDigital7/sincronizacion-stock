from __future__ import annotations

from io import BytesIO

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def build_excel(summary: pd.DataFrame, detail: pd.DataFrame) -> bytes:
    output = BytesIO()
    mismatches = detail[detail["estado_sincronizacion"] != "SINCRONIZADO"]
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Resumen", index=False)
        mismatches.to_excel(writer, sheet_name="Detalle", index=False)
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for idx, column in enumerate(sheet.columns, 1):
                width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
                sheet.column_dimensions[get_column_letter(idx)].width = width
    return output.getvalue()

